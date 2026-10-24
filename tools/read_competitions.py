# -*- coding: utf-8 -*-
"""读取竞赛赛事列表 docx，导出为 JSON。"""
from docx import Document
import json
from pathlib import Path

SRC = Path(r"C:\Users\lenovo\Downloads\附件：合肥工业大学学科竞赛赛事列表（2026年版）.docx")
OUT = Path(r"C:\Users\lenovo\Documents\deepseek-harness\default-workspace\campusknow\site\data\competitions.json")

doc = Document(str(SRC))
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))

rows_out = []
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        # 去重相邻单元格（合并单元格会出现重复文本）
        merged = []
        for c in cells:
            if not merged or merged[-1] != c:
                merged.append(c)
        if any(merged):
            rows_out.append({"table": ti, "row": ri, "cells": merged})

print("data rows:", len(rows_out))
for r in rows_out[:6]:
    print(" | ".join(r["cells"])[:120])

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps({"source": SRC.name, "rows": rows_out},
                          ensure_ascii=False, indent=1), encoding="utf-8")
print("saved ->", OUT)
